"""Rebuild Nova's evaluation artifacts using existing python-pptx and Pillow.

Only writes beneath docs/hackathon-presentation and docs/architecture.
No application imports, private database reads, live audio, or network requests.
"""
from pathlib import Path
import ast
import hashlib
import json
import shutil
import textwrap
import xml.etree.ElementTree as ET
from zipfile import ZipFile
from html import escape
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / 'docs/hackathon-presentation'
ARCH = ROOT / 'docs/architecture'
OUT.mkdir(parents=True, exist_ok=True)
ARCH.mkdir(parents=True, exist_ok=True)
DATE = '10 October 2026'
NAVY = '#10263A'
TEAL = '#007F86'
INK = '#18384C'
MUTED = '#526879'
PALE = '#EAF3F4'
AMBER = '#A35B0B'

def write(name, content):
    # Expand compact prose measurements without touching repository identifiers.
    replacements = {'median1,':'median 1,', 'p951,':'p95 1,', 'median871':'median 871',
        'p953,':'p95 3,', 'median2,':'median 2,', 'was2,':'was 2,', 'median3.':'median 3.',
        'p956.':'p95 6.', 'over100':'over 100', 'and303':'and 303', 'reports17':'reports 17',
        'in61.':'in 61.', 'seconds,187':'seconds, 187', 'umask0077':'umask 0077',
        'exit78':'exit 78', 'exit75':'exit 75', 'queue50':'queue 50', 'queue2':'queue 2',
        'threshold0.018':'threshold 0.018', 'Python3.13':'Python 3.13', 'Python3.12':'Python 3.12',
        'Pi5':'Pi 5', 'Llama3.2':'Llama 3.2', 'about2.48':'about 2.48', 'and1,024':'and 1,024',
        'reports1,278':'reports 1,278', 'gpiozero2':'gpiozero 2', 'at least30':'at least 30',
        'speaker output and16kHz':'speaker output and 16 kHz', 'Linux64-bit':'Linux 64-bit',
        'BCM17/27/22':'BCM17 / 27 / 22', 'chip0':'chip 0', 'retention0':'retention 0'}
    for old,new in replacements.items(): content=content.replace(old,new)
    (OUT / name).write_text(content.strip() + '\n', encoding='utf-8')

def ref(path, needle=None):
    lines = (ROOT / path).read_text(encoding='utf-8-sig', errors='replace').splitlines()
    n = next((i for i, line in enumerate(lines, 1) if needle and needle in line), 1)
    return f'{path}:{n}'

# Inventory all first-party text files, not private .env/data or bundled vendors.
inventory = []
for folder in ['src', 'tests', 'scripts', 'config', 'docs', 'deployment/raspberry-pi', 'deployment/windows']:
    for p in sorted((ROOT / folder).rglob('*')):
        relative = p.relative_to(ROOT).as_posix()
        if not p.is_file() or relative.startswith(('docs/hackathon-presentation/', 'docs/architecture/')):
            continue
        if any(x in p.parts for x in ['__pycache__', 'espeak-ng-data', 'whisper.cpp', 'ollama']):
            continue
        if p.suffix not in ['.py', '.md', '.json', '.sh', '.txt', '.toml', '.in', '.example']:
            continue
        raw = p.read_bytes()
        source = raw.decode('utf-8-sig', errors='replace')
        entry = {'path': relative, 'bytes': len(raw), 'lines': len(source.splitlines()),
                 'sha256': hashlib.sha256(raw).hexdigest()}
        if p.suffix == '.py':
            tree = ast.parse(source)
            entry['classes'] = [n.name for n in ast.walk(tree) if isinstance(n, ast.ClassDef)]
            entry['functions'] = [{'name': n.name, 'line': n.lineno} for n in ast.walk(tree)
                                  if isinstance(n, (ast.FunctionDef, ast.AsyncFunctionDef))]
            entry['imports'] = sorted({n.module or '' for n in ast.walk(tree) if isinstance(n, ast.ImportFrom)} |
                                      {a.name for n in ast.walk(tree) if isinstance(n, ast.Import) for a in n.names})
            entry['syntax'] = 'parsed successfully'
        elif p.suffix == '.md':
            entry['headings'] = [x for x in source.splitlines() if x.startswith('#')]
        inventory.append(entry)
for name in ['README.md', 'SECURITY.md', 'LICENSE.md', 'pyproject.toml', 'requirements.txt', '.gitignore']:
    p = ROOT / name
    raw = p.read_bytes()
    inventory.append({'path': name, 'bytes': len(raw), 'lines': len(raw.splitlines()),
                      'sha256': hashlib.sha256(raw).hexdigest()})
write('source-inventory.json', json.dumps(inventory, indent=2))

FEATURES = [
 ('Continuous speech capture', 'Implemented', 'audio/continuous_recorder.py', 'Repeated utterance capture, not uninterrupted recording during inference/playback.'),
 ('VAD', 'Implemented', 'audio/vad.py', 'RMS energy threshold, not a neural VAD.'),
 ('Whisper STT', 'Implemented', 'stt/whisper.py', 'Local whisper.cpp executable and pre-provisioned model.'),
 ('Nova wake activation', 'Implemented', 'core/wake_detector.py; core/activation_controller.py', 'Transcript-prefix matching after STT, one 8-second follow-up after wake-only speech.'),
 ('Local Llama chat', 'Implemented', 'llm/ollama.py; llm/local_endpoint.py', 'Llama 3.2 1B default, non-streaming loopback HTTP, not a bundled trained model.'),
 ('SQLite facts, notes, tasks, reminders', 'Implemented', 'memory/database.py; memory/repository.py; tools/memory_tools.py', 'Plaintext local persistence, transactional operations.'),
 ('Passive memory', 'Implemented / opt-in', 'memory/passive_worker.py; memory/passive_processor.py', 'Bounded asynchronous queue and grounded observations. Disabled by default.'),
 ('Semantic recall', 'Implemented / limited', 'memory/hybrid_retriever.py', 'Small concept lexicon and cosine ranking. No embeddings or unrestricted semantic search.'),
 ('Review, correction, deletion', 'Implemented', 'memory/management.py; tools/lifecycle.py', 'Scoped confirmations, typed IDs, transaction rechecks. No speaker authentication.'),
 ('Reminder notifications', 'Implemented', 'scheduler/reminder_scheduler.py; scheduler/reminder_notifications.py', 'Leases, acknowledgment, retries. Crash window can duplicate audible delivery.'),
 ('Piper TTS and playback', 'Implemented', 'tts/piper.py; audio/player.py', 'Local synthesis and sounddevice output. Normal capture pauses for playback.'),
 ('Software mute and privacy gates', 'Implemented', 'core/privacy.py; core/privacy_policy.py; audio/temporary.py', 'Generation guards, secret heuristics, owned WAV cleanup. No encryption/secure erasure.'),
 ('RGB LED and physical switch', 'Implemented / hardware unverified', 'core/hardware.py', 'gpiozero/LGPIOFactory backend, mock/no-op backends. GPIO switch is a software permission input.'),
 ('Pi 5 Linux ARM64 and systemd', 'Implemented preparation / target unverified', 'core/pi_deployment.py; core/systemd_units.py; deployment/raspberry-pi/', 'Provisioning, guarded launch, non-root service templates. No proven physical Pi run.'),
 ('Weather and news', 'Implemented / opt-in / live unverified', 'tools/online_tools.py; tools/online_requests.py', 'Open-Meteo and NewsAPI only. Internet and provider setup required.'),
 ('Spoken interruption', 'Partial / experimental', 'audio/stop_monitor.py; core/cancellation.py', 'Exact wake+stop monitor, recognition gaps, no AEC or general full duplex.'),
 ('Packaged voice entry point', 'Partial', 'main.py; pyproject.toml', 'Installed CLI prints startup text. Voice loop lives in scripts/always_on_assistant.py.'),
 ('Multi-hour reliability', 'Deferred validation', 'docs/LONG_RUNNING_VALIDATION.md', 'Short synthetic soak exists. Multi-hour physical soak remains deferred.'),
 ('Encryption, voice identity, electrical mic cutoff', 'Future improvement / absent', 'No corresponding implementation found', 'Do not present these as existing capabilities.'),
 ('Cloud sync, mobile app, unrestricted agents', 'Absent', 'No corresponding implementation found', 'No arbitrary shell/IoT/calendar/email tool execution layer found.'),
]

report = f'''# Nova repository evaluation

Audit date: {DATE}. Scope: first-party Python source, tests, configuration examples,
deployment scripts, service templates and documentation. The inventory records hashes,
line counts and Python symbols for reproducibility. Private `.env`, personal SQLite
records and recordings were deliberately excluded. Bundled binaries/model assets were
inventoried by role and presence, not reverse-engineered or treated as proof of execution.

## Evaluation position

Nova is a substantial local voice-assistant prototype with persistent personal memory,
deterministic task/reminder tools, guarded passive observations and privacy controls.
Its strongest claim is integration and evidence-aware memory on a local device.
There is no evidence of a newly trained foundation model or novel acoustic wake engine.
Raspberry Pi 5 is the documented deployment target, not a demonstrated working device
in this audit. Do not equate passing mocked tests with hardware readiness.

## Implemented, partial and future capabilities

Package-relative paths below are under `src/offline_ai_companion/`, except explicit paths.

| Capability | Status | Evidence | Qualification |
| --- | --- | --- | --- |
'''
for a,b,c,d in FEATURES:
    report += f'| {a} | {b} | `{c}` | {d} |\n'
report += '''
## Actual architecture

### Runtime and entry points

`scripts/always_on_assistant.py` constructs recorder, player, STT, Assistant, Piper,
activation controller, privacy gate, optional passive worker and reminder scheduler.
Its main path delegates to `core/service_runtime.py` for console/GPIO controls and
SIGTERM handling. `main.py` only prints startup messages, despite the packaged
`offline-ai-companion` entry point. Use the script or guarded Pi launcher for a demo.

### Audio ingestion and voice activation

The recorder opens sounddevice/PortAudio mono float32 capture at 16 kHz in 100 ms
blocks. The RMS threshold defaults to 0.018, adjustable through `VAD_THRESHOLD`.
Defaults include 0.7 seconds of pre-roll, 0.2 seconds of consecutive speech to start,
1 second of silence to finish and a 15-second utterance maximum with a sample-count
bound. A bounded queue avoids unlimited audio accumulation. Dropped blocks or input
overflow discard the utterance. Five seconds without blocks raises a recoverable error.

Utterances become PCM16 WAVs in an owned session directory. WhisperSTT invokes a local
whisper.cpp CLI with the local tiny.en model by platform default and an English Nova
prompt. It validates files and bounds subprocess lifetime. Caption/invalid-transcript
filters run before dispatch. WakeDetector matches `Nova` or `hey Nova` at the beginning
of the transcript using regex boundaries. STT runs BEFORE wake filtering. Nova therefore
transcribes VAD-triggered speech even with passive storage disabled. This is not a low
power acoustic keyword detector.

A wake+command utterance routes immediately. A wake-only utterance opens an eight-second
window for one subsequent utterance, using timezone-aware capture timestamps, then
resets. Other speech never receives an active spoken answer. If passive memory is
enabled, eligible transcripts enter its bounded worker queue, otherwise they are dropped.
Always-on means a repeating listening loop while available: the normal path serially
records, transcribes, processes and speaks. It cannot capture every utterance during
ordinary processing or output. Half-duplex playback invalidates old recordings and
waits for capture closure, with a 0.3-second acoustic settling delay.

### Routing, local model and memory

Assistant first enforces credential policy and pending management/confirmation workflows.
Grounded exact/hybrid recall can answer without model generation. QueryRouter classifies
supported deterministic tools and general chat. Personal fact parsing proposes changes
and typically asks for confirmation. The legacy name update, explicit note/task creation
and reminder creation are direct writes, so “every write requires confirmation” is false.
LLM-assisted fact extraction creates proposals rather than trusting arbitrary output.

OllamaLLM posts to `/api/generate` using loopback-only HTTP(S), default model
`llama3.2:1b`, `stream=false`, temperature 0.4 and 160 predicted tokens. It validates
response format and limits retries (normally one retry, 120-second timeout). Runtime
policy rejects remote endpoints, redirects and environment proxies. In experimental
interruption mode, active requests use a 30-second timeout and zero retries. Cancellation
can suppress results but cannot guarantee that Ollama's server-side inference stops.

SQLite tables hold facts, notes, tasks, reminders and passive memories. Passive records
carry category, subject, content, evidence, capture time, optional event time and unique
fingerprint. Connections enable foreign keys, commit on success, roll back on failure
and close. Lifecycle migrations add revision and delivery-lease fields without resetting
IDs. Conversation history is RAM-only, bounded to ten turns and 4,000 context characters.
The constructed general-chat memory context includes all confirmed facts and capped
notes/tasks/reminders/passive rows, so large confirmed-fact sets can exceed a small model
context. This is a remaining scaling risk, not an implemented full token-budget manager.

Passive extraction first uses supported deterministic assertion patterns, otherwise local
LLM JSON candidates. Validation checks verbatim evidence, assertion/negation, supported
subject/category/content and grounded times before saving. Inputs above 1,500 characters
are rejected by the processor, with at most five items of up to 300 content characters.
The worker is asynchronous and non-daemon, rejects full queues and drains accepted work
on shutdown. Optional collection defaults OFF. Observations belong to an unidentified
speaker, never automatically to the owner. They never become actionable tasks/reminders.

Hybrid recall uses exact relations first and a small synonym/concept lexicon plus sparse
cosine ranking. It examines at most 32 of 33 fetched candidates, returns at most three
ranked observations, and uses short SQLite lock/progress budgets. Confirmed facts take
priority. Evidence and selected-record rechecks prevent similarity alone from establishing
truth. There is no embedding/vector database, arbitrary semantic reasoning, speaker
diarization or authentication. Unknown concepts and date relations have conservative limits.

### Tools and scheduling

Local tools add/review notes, tasks and reminders, review/correct/delete facts/passive
records and complete/rename/delete tasks or cancel/snooze/reschedule reminders. Management
uses typed IDs, expiring RAM-only confirmations (120 seconds by default), transaction
rechecks and revisions to resist concurrent changes. Creation has the direct-write
exceptions noted above. Reminder date parsing uses an IANA timezone, defaults to
Asia/Kolkata, persists UTC and rejects unsupported/ambiguous or past calendar times.

The scheduler polls every second by default, atomically claims due rows with renewable
leases and requests playback through the main loop. Notifications interrupt capture.
Successful playback acknowledgment precedes completion persistence. Playback or commit
failure leaves recovery work, and interruption defers a reminder by 60 seconds. A crash
after audible playback and before durable completion can cause repeat playback: delivery
is not exactly once. Reminder speech can still occur while microphone input is muted.

Optional weather/news intents use fixed HTTPS Open-Meteo geocoding/forecast and NewsAPI
endpoints. Global enablement defaults false. Only explicit/configured locations and
configured news topics go to providers, with the NewsAPI key for authentication.
Memory/history do not enter provider requests. Sanitized provider answers bypass Ollama
and persistent memory. Weather is modeled current/daily information, not a local sensor
or offline forecast. These are the only external application integrations found.

### Output, hardware and operating system

PiperTTS sends text to a local executable using an argument list and stdin, producing a
WAV from `en_US-lessac-medium.onnx` by default. soundfile/sounddevice perform playback.
Experimental playback checks cancellation between 20 ms writes. The normal path has
no acoustic echo cancellation or full duplex. The optional exact wake+stop monitor
alternates recording/STT and discards other speech, leaving recognition gaps.

The GPIO module explicitly uses local gpiozero DigitalInput/OutputDevice with LGPIOFactory,
plus pure-Python mock and fail-closed no-op backends. Defaults are BCM17 red, BCM27 green,
BCM22 blue and BCM23 switch, chip0, active-high LED and muted-high switch. Red means muted,
green listening, blue processing/speaking, off idle/interrupted/shutdown. It polls at
20 ms and debounces release for 100 ms. Every GPIO-controlled restart starts muted and
requires a stable muted-to-listen switch cycle. Any GPIO read/write fault latches mute
until restart. The switch asserts software permission; it does not electrically remove
microphone power or guarantee zero-latency isolation.

The Pi profile requires Linux ARM64 and Python 3.13. Its conservative defaults are tiny.en,
Llama 3.2 1B, two Whisper/Ollama threads, 1,024 model context, audio queue50 and passive
queue2. It starts muted and disables passive extraction, interruption and online tools.
Provisioning scripts install OS/Python dependencies, build pinned whisper.cpp v1.7.6,
download voices/models explicitly and install a specified Ollama archive with a supplied
SHA256. They do not establish target compatibility by themselves. Generated model hashes
aid reproduction but are not independent trusted pre-download digests for every model.

The hardware script exposes distro gpiozero/lgpio in the venv and adds audio/gpio groups.
The service installer renders units, mandates target `systemd-analyze verify`, installs
without enabling/starting and uses a non-root service user. Nova uses umask0077,
NoNewPrivileges and ProtectSystem=full. SIGTERM follows graceful cleanup, rejects new
work and drains accepted passive items. Restart policy limits three failures per five
minutes, delays ten seconds and suppresses restart for configuration exit78. Transient
preflight readiness uses exit75. Nova's infinite stop timeout preserves draining but
can leave a stalled service requiring operator intervention. Physical boot, GPIO,
headless ALSA, ARM64 inference and thermal capacity remain unverified.

## Offline boundary and privacy

Core STT, model inference, memory, tools and TTS can operate without Internet after
local dependencies/models are provisioned. Offline capability is configuration-dependent:
online tools explicitly cross the boundary when enabled, and provisioning needs Internet.
Runtime local-model enforcement is an application policy, not a firewall or a guarantee
about independently configured Ollama/binary logs, telemetry or forwarding.

Software/physical mute combine with OR semantics. Generation/cancellation guards invalidate
stale input and stop new capture/dispatch. Work admitted earlier can finish and accepted
passive work drains. Routine runtime logs use fixed labels, counts and error categories.
Private debug/manual scripts can expose content and retain recordings. Owned runtime WAVs
are cleaned on normal/error shutdown, but power loss/SIGKILL can leave files and there
is no automatic crash sweep. SQLite is plaintext. Retention defaults to zero (no age
expiry), passive expiry hides records during reads, and explicit purging removes rows.
Deletion is not secure erasure. Regex credential detection is useful but incomplete.
No account authorization or speaker identity protects voice operations from nearby people.

## Evidence and performance claims

The latest DEVELOPMENT_STATUS checkpoint reports 1,278 tests passing in 114.53 seconds.
This is historical evidence, not automatically a result of this audit. See verification.md
and test-results.xml for the new attempt and runtime qualification.

The 9 October file-input benchmark in module-1266-validation-results.md reports Whisper
median1,133.163 ms/p951,278.519 ms (10 samples), Piper median871.319 ms/p951,968.337 ms
(15 mixed samples), and a file-response median2,481.135 ms/p953,132.096 ms (5 samples).
The file-response benchmark excludes microphone/VAD, physical playback and settling.
Its single Ollama probe was2,738.603 ms. Do not add these medians or advertise them as a
Pi end-to-end latency or accuracy distribution. Expanded retrieval's synthetic benchmark
reports median3.409 ms/p956.585 ms over100 queries and303 records, not an archive-scale test.
The synthetic soak JSON reports17 cycles in61.34 seconds,187 stream opens/closes and
one thread before/after, with fake models and no physical devices. None proves months
of uptime, thermal safety or multi-hour capture.

The newest status document says the user later reported Windows physical audio and
interruption checks passed. Older module checklists still say NOT RUN. Treat this as
reported manual evidence with stale snapshots, not a new independently observed run.

## Evaluation gaps and demo risks

1. **High: inaccessible project Python in this environment.** The Store Python3.13/venv
   launcher fails to execute here. Restore/verify a supported interpreter before the
   live demonstration. Bundled Python3.12 can author files and offers a qualified test
   attempt using existing packages, but is outside the declared >=3.13,<3.14 runtime.
2. **High: Pi hardware success unproven.** Keep the working Windows device as the primary
   demo unless the target Pi passes physical checks. Label mock GPIO honestly.
3. **High: privacy is software enforced.** No voice identity, encryption or electrical
   microphone cutoff. Obtain participant consent for any passive collection demo.
4. **Medium: STT/VAD venue noise and sequential capture.** A loud hall can trigger false
   utterances. Rehearse mic placement and threshold, use short sentences and wait for
   LISTENING. A dedicated acoustic wake engine is future work.
5. **Medium: performance/resource uncertainty.** LLM cold starts, background inference,
   1,024-token Pi context, CPU contention and thermal throttling need target measurement.
6. **Medium: service/audio dependencies.** Python3.13 requirement, ARM64 wheel availability,
   RP1 chip mapping and headless audio devices must be checked on the target.
7. **Medium: reminders and interruption.** Verify audible delivery and restart recovery.
   Explain the duplicate crash window. Keep experimental spoken stop out of the core demo.
8. **Medium: public reproducibility.** Root requirements.txt embeds an absolute Windows
   editable path. pyproject omits audio extras. Use platform setup guides rather than
   treating root pip freeze as portable. No CI workflow or general release packaging found.
9. **Medium: documentation completeness.** README.md, SECURITY.md, LICENSE.md and
   docs/architecture.md, development.md, setup.md and raspberry-pi.md are empty files.
   The actual Pi guide is deployment/raspberry-pi/README.md. Historical docs conflict.
   Complete public setup, licensing and security documentation before distribution.

## Recommended next work (not implemented claims)

First validate a real Pi install, GPIO fail-closed behavior, boot/restart and multi-hour
audio/model workload with temperature, RSS, CPU and p50/p95 end-to-end timing. Then improve
deployment reproducibility and package the real voice launcher. Add a dedicated acoustic
wake detector and better VAD after measuring missed/false activation rates. Define consent,
retention and threat-model policy, consider encrypted storage and a true physical mic
disconnect. Add authentication before privileged integrations. Bound the entire memory
prompt by tokens, benchmark larger archives, and evaluate embedding retrieval only if
memory/CPU budgets allow it. AEC/full-duplex, cloud sync and mobile interfaces remain future
projects requiring separate implementation and validation.

## Source navigation

The exported source-inventory.json lists all scanned first-party files and symbols.
Key anchors are generated below from this exact checkout. Source hashes allow reviewers
to distinguish this audit snapshot from later source changes.
'''
anchors = [ref('scripts/always_on_assistant.py', 'class AlwaysOnAssistant'),
 ref('src/offline_ai_companion/audio/continuous_recorder.py', 'class ContinuousRecorder'),
 ref('src/offline_ai_companion/audio/vad.py', 'class VoiceActivityDetector'),
 ref('src/offline_ai_companion/stt/whisper.py', 'class WhisperSTT'),
 ref('src/offline_ai_companion/core/wake_detector.py', 'class WakeDetector'),
 ref('src/offline_ai_companion/core/assistant.py', 'class Assistant'),
 ref('src/offline_ai_companion/core/hardware.py', 'class HardwareController'),
 ref('src/offline_ai_companion/core/pi_deployment.py', 'def preflight'),
 ref('src/offline_ai_companion/core/systemd_units.py'),
 ref('src/offline_ai_companion/tools/online_tools.py'),
 ref('src/offline_ai_companion/memory/hybrid_retriever.py'),
 ref('src/offline_ai_companion/core/privacy_policy.py'),
 ref('src/offline_ai_companion/main.py', 'def main')]
report += '\n'.join(f'- `{a}`' for a in anchors)
write('repository-evaluation.md', report)

QA = [
 ('Innovation', 'What is actually new about Nova?', 'The integration of local speech, persistent memory, grounded passive observations and deterministic tools with shared mute controls. It reuses existing models. No new trained model or acoustic wake algorithm is claimed.', 'core/assistant.py; memory/passive_processor.py'),
 ('Innovation', 'Why use a small local model?', 'Llama3.2:1b is the configured default and fits the project’s constrained-device intent. The repository provides thread/context bounds, but does not prove Pi latency or RAM sufficiency.', 'config.py; deployment/raspberry-pi/nova.env.example'),
 ('Architecture', 'Describe the complete voice path.', 'Mic capture and RMS VAD produce a WAV. whisper.cpp transcribes it. Activation routes the text to tools/recall/chat or optional passive storage. Piper synthesizes an active answer and sounddevice plays it.', 'scripts/always_on_assistant.py'),
 ('Architecture', 'Is Nova a real acoustic wake-word detector?', 'No. WakeDetector regex-matches the beginning of Whisper’s transcript. The recorder and STT precede wake filtering, so it is not a low-power keyword engine.', 'core/wake_detector.py; stt/whisper.py'),
 ('Architecture', 'What happens if I say only Nova?', 'It opens an eight-second window for one follow-up utterance. A command in the same utterance executes directly. Explicit wake speech takes precedence and mute resets activation.', 'core/activation_controller.py; scripts/always_on_assistant.py'),
 ('Architecture', 'Does always-on mean continuous recording without gaps?', 'No. The normal voice loop is sequential and half-duplex. Listening resumes after processing/output. Input closes for playback and a settling interval.', 'audio/half_duplex.py; scripts/always_on_assistant.py'),
 ('Architecture', 'Which VAD model do you use?', 'No neural VAD. RMS energy is compared with threshold0.018 by default. Continuous capture adds pre-roll, minimum speech and silence/maximum-duration bounds.', 'audio/vad.py; audio/continuous_recorder.py'),
 ('Offline', 'Does it work with Internet disconnected?', 'The core can, after models/dependencies are provisioned locally and Ollama runs locally. Weather/news need Internet and are disabled by default. We still require a physical disconnection rehearsal.', 'llm/local_endpoint.py; tools/online_tools.py'),
 ('Offline', 'Why is there HTTP in an offline assistant?', 'Ollama is a separate local process accessed over loopback HTTP. This is local IPC through an API, not a cloud request. Remote URLs, redirects and environment proxies are blocked.', 'llm/local_endpoint.py; llm/ollama.py'),
 ('Offline', 'Does setup itself require Internet?', 'Yes, explicit OS/package/model provisioning does. Runtime does not automatically download models. Offline operation starts after provisioning.', 'deployment/raspberry-pi/setup.sh; models.sh'),
 ('Privacy', 'What is stored when nobody says Nova?', 'By default passive memory is disabled, so recognized background text is not queued for memory. If enabled, eligible speech can become grounded local observations without a spoken response.', 'config.py; memory/passive_worker.py'),
 ('Privacy', 'Does disabling passive memory stop transcription?', 'No. STT precedes transcript wake detection. Use software or physical mute to deny new capture and dispatch.', 'scripts/always_on_assistant.py; core/privacy.py'),
 ('Privacy', 'Does the physical switch cut microphone power?', 'No. GPIO informs the software privacy gate. Polling and cancellation cannot establish zero-latency electrical isolation. A power/data disconnect would need additional hardware.', 'core/hardware.py'),
 ('Privacy', 'What happens to queued work after mute?', 'New/stale input is rejected by generations and guards. Work already accepted may finish and passive work drains. Committed facts/tasks/reminders remain.', 'core/privacy.py; memory/passive_worker.py'),
 ('Privacy', 'Are recordings immediately and permanently destroyed?', 'Owned runtime WAVs are normally removed after their consumer finishes or on cleanup. Abrupt termination can leave files. Removal is not secure erasure, and manual scripts retain sample recordings.', 'audio/temporary.py; docs/PRIVACY.md'),
 ('Privacy', 'Is memory encrypted?', 'No. SQLite is plaintext. Private file modes/umask and account protection help, but encryption and secure deletion are absent.', 'memory/database.py; deployment/raspberry-pi/systemd/nova.service.in'),
 ('Memory', 'How do you distinguish memory from hallucination?', 'Confirmed facts outrank observations. Passive candidates need supported assertion evidence, subject/category and time grounding. Recall rechecks evidence and saved rows. General model chat can still hallucinate.', 'memory/passive_processor.py; memory/hybrid_retriever.py'),
 ('Memory', 'Do passive observations belong to the device owner?', 'No. They are attributed to an unidentified speaker. There is no diarization or voice authentication, and recall states this uncertainty.', 'memory/passive_retriever.py; memory/hybrid_retriever.py'),
 ('Memory', 'Is semantic recall based on embeddings?', 'No. It uses a small concept/synonym lexicon with sparse cosine ranking over bounded SQLite candidates after exact relation checks.', 'memory/hybrid_retriever.py'),
 ('Memory', 'Can it recall arbitrary paraphrases?', 'Coverage is limited by the concept vocabulary and evidence rules. Unknown relations, qualifiers and dates can return unavailable rather than a guessed personal fact.', 'memory/hybrid_retriever.py; docs/SEMANTIC_MEMORY.md'),
 ('Memory', 'Does every memory write need confirmation?', 'No. Proposed ordinary facts and management changes have confirmations, but legacy name setting, explicit notes/tasks and reminder creation can write directly. Passive observations have policy checks rather than voice confirmation.', 'core/assistant.py; tools/memory_tools.py; tools/lifecycle.py'),
 ('Memory', 'Can I correct or delete saved memory?', 'Yes, scoped commands review facts/passive records and propose corrections/deletions with expiring confirmations and transactional rechecks. Deleting a fact preserves separate passive history.', 'memory/management.py'),
 ('Memory', 'Does retention actually delete old rows?', 'Read-time passive expiry hides old observations. Explicit purge removes them. Default retention0 means no age expiry. Confirmed memory is not automatically deleted.', 'memory/passive_repository.py; config.py'),
 ('Tools', 'Can background speech schedule a reminder?', 'No. Passive task/event records are observations. Only explicit active tool commands create actionable tasks and reminders.', 'memory/passive_processor.py; tools/lifecycle.py'),
 ('Tools', 'Can the LLM run shell commands?', 'No arbitrary shell-agent tool exists. QueryRouter selects supported deterministic memory/lifecycle and opt-in weather/news handlers.', 'router/query_router.py; core/assistant.py'),
 ('Tools', 'What are the online integrations?', 'Open-Meteo geocoding/current modeled weather/daily forecast and NewsAPI headlines. They require explicit enablement/provider configuration and live validation is unperformed.', 'tools/online_tools.py; docs/ONLINE_TOOLS.md'),
 ('Security', 'What personal information leaves for weather/news?', 'Explicit/configured location or allowed configured news topic and provider authentication as needed. The tools do not send memory/history, and answers bypass Ollama and memory storage.', 'tools/online_requests.py; tools/online_tools.py'),
 ('Security', 'How do you protect against prompt injection?', 'Deterministic tools do not trust model instructions, passive output is grounded, memory is marked data in chat prompts, and provider answers bypass the LLM. These controls do not make general chat immune to injection.', 'core/assistant.py; memory/passive_processor.py; tools/online_tools.py'),
 ('Security', 'Can Nova store passwords?', 'Recognized credential patterns are refused at several boundaries and omitted from context/output. Regex detection is heuristic, so it is not a credential vault or complete data-loss prevention system.', 'core/privacy_policy.py; memory/repository.py'),
 ('Security', 'Who can issue destructive voice commands?', 'Anyone with access to the active voice/text interface can issue and confirm supported changes. Confirmations guard intent/concurrency, not identity. Authentication is future work.', 'memory/management.py; tools/lifecycle.py'),
 ('Reliability', 'How are reminder duplicates prevented?', 'SQLite claims and renewable leases avoid competing active deliveries. Successful playback must be acknowledged before completion. A crash between sound and durable completion can still repeat delivery.', 'scheduler/reminder_scheduler.py; memory/repository.py'),
 ('Reliability', 'Does microphone mute also silence reminders?', 'No general silent-output mode is implemented. The loop checks pending reminders before checking input mute, so microphone privacy does not promise speaker silence.', 'scripts/always_on_assistant.py'),
 ('Reliability', 'What if a device stops sending audio?', 'After the stall timeout, capture raises a recoverable error. Recovery logs a sanitized category, waits and retries. Overflowed or dropped-block utterances are discarded.', 'audio/continuous_recorder.py; audio/recovery.py'),
 ('Reliability', 'Can you interrupt a response?', 'Experimental console cancellation and optional exact wake+stop monitoring exist. Child STT/TTS and owned playback support cancellation, but Ollama inference may continue and the monitor has recognition gaps.', 'core/cancellation.py; audio/stop_monitor.py'),
 ('Reliability', 'Is full-duplex conversation implemented?', 'No. Half-duplex is the default, and the experimental stop monitor is narrower than general barge-in. No acoustic echo cancellation exists.', 'audio/half_duplex.py; audio/stop_monitor.py'),
 ('Performance', 'What latency can you substantiate?', 'Historical Windows file-input measurements report about2.48 seconds median for five recall responses without capture/playback. They are not full physical latency, Pi measurements or fresh results from this audit.', 'docs/module-1266-validation-results.md'),
 ('Performance', 'How much memory and CPU does Pi need?', 'The profile targets constrained resources with tiny.en, 1B LLM, two threads and1,024 context. Actual Pi RAM, CPU, temperature and end-to-end timing remain unmeasured here.', 'deployment/raspberry-pi/nova.env.example; README.md'),
 ('Scalability', 'What happens with a large memory archive?', 'Recall bounds candidates/results and lock waits, but concept coverage is limited. General-chat context includes all facts, so total prompt size can grow. Large-archive and multi-user performance are not validated.', 'memory/hybrid_retriever.py; core/assistant.py'),
 ('Scalability', 'Can several users share the companion?', 'There is no user-account isolation or voice identity. The current design is a single local companion; multi-user permissions and separate memory namespaces need new work.', 'memory/database.py; core/assistant.py'),
 ('Deployment', 'Have you actually deployed to Pi5?', 'The repository has ARM64 provisioning, preflight, local GPIO and systemd code. Documentation explicitly says physical GPIO, Linux audio, boot/recovery and ARM64 inference have not been validated.', 'docs/DEVELOPMENT_STATUS.md; deployment/raspberry-pi/README.md'),
 ('Deployment', 'How does it start after reboot?', 'Installed systemd units can be explicitly enabled. Nova starts muted and a stable physical muted-to-listen switch cycle authorizes capture. Installation alone does not enable services.', 'core/hardware.py; deployment/raspberry-pi/install-service.sh'),
 ('Deployment', 'What happens if GPIO breaks?', 'Read/write faults latch software and physical mute until restart. No automatic recovery unmute occurs. An LED fault can prevent a reliable red/off indication.', 'core/hardware.py'),
 ('Deployment', 'Why can shutdown take a long time?', 'Accepted passive work drains, and Nova’s systemd stop timeout is infinite to avoid killing it. A stalled shutdown needs operator diagnosis. This is an explicit reliability tradeoff.', 'memory/passive_worker.py; deployment/raspberry-pi/systemd/nova.service.in'),
 ('Evaluation', 'What do the tests prove?', 'They cover routing, persistence, privacy gates, failure recovery, lifecycle and mocked voice/GPIO/service policies. The historical latest checkpoint reports1,278 passing tests. They do not prove physical Pi operation.', 'tests/; docs/DEVELOPMENT_STATUS.md'),
 ('Evaluation', 'What is your biggest current limitation?', 'The target Pi and sustained physical workload remain unvalidated, and sequential audio/STT-based wake handling impose gaps and cost. Privacy also lacks identity, encryption and electrical isolation.', 'docs/DEVELOPMENT_STATUS.md; core/hardware.py'),
 ('Future', 'What would you improve first?', 'Validate the physical Pi and long-running workload, repair deployment/entry-point/documentation gaps, then measure acoustic wake/VAD improvements and stronger consent, authentication and storage controls.', 'repository-evaluation.md recommendations'),
 ('Future', 'Would a larger model solve everything?', 'No. It may improve language quality but increases local resource demand and does not fix acoustic capture, identity, electrical mute, database privacy or service validation.', 'config.py; core/hardware.py; memory/database.py'),
 ('Demo', 'How will you show that memory is persistent?', 'Use an isolated demo SQLite path, save a synthetic fact/note, restart with the same path and recall it. Never demonstrate against the personal production database.', 'memory/database.py; scripts/always_on_assistant.py'),
]
qa = '# Hackathon jury cross-questions\n\n' + f'{len(QA)} repository-grounded questions. Package paths below are under `src/offline_ai_companion/` unless explicit. Hardware and benchmark qualifications are deliberate.\n'
for i, (category, question, answer, evidence) in enumerate(QA, 1):
    qa += f'\n## {i}. [{category}] {question}\n\n{answer}\n\nEvidence: `{evidence}`.\n'
write('jury-cross-questions.md', qa)

write('demo-script.md', '''# Nova jury demonstration script

Target duration: 6 minutes. Rehearse on the actual presentation machine. This script is
a proposed demonstration, not evidence that the physical run happened in this audit.

## Before the jury enters

Use synthetic speech only. Keep the same isolated database for restart demonstration.
Models and voice must already exist, Ollama must already serve llama3.2:1b on loopback,
and the supported Python3.13 interpreter must work. The current project Store/venv
launcher is inaccessible in the audit sandbox, so resolve that before presentation.
Use the metadata-only scripts/audio_diagnostics.py for format/device preflight.
Do not run an unreviewed health check against the default personal database: some
health-check paths initialize/write SQLite. Any database check must use the demo path.

In a NEW PowerShell terminal from the repository root, after fixing the interpreter:

```powershell
$env:DATABASE_PATH = Join-Path $PWD 'docs/hackathon-presentation/.build/demo-companion.db'
$env:PASSIVE_MEMORY_ENABLED = 'false'
$env:NOVA_ONLINE_TOOLS_ENABLED = 'false'
$env:NOVA_EXPERIMENTAL_INTERRUPTION = 'false'
$env:NOVA_SPOKEN_INTERRUPTION = 'false'
$env:NOVA_PRIVATE_DEBUG = 'false'
$env:MIC_ENABLED = 'true'
$env:NOVA_GPIO_BACKEND = 'none'
$env:NOVA_HEADLESS = 'false'
$env:NOVA_PERFORMANCE_LOG = 'true'
.\.venv\Scripts\python.exe .\scripts\always_on_assistant.py
```

Do not use the installed offline-ai-companion command for the voice demo: it currently
only prints startup text. Use the actual script above. A separate shell keeps overrides
out of ordinary use. Close it afterward. The demo database is retained intentionally.
For a Pi, use its separate env/profile and guarded launcher after target preflight.
Never copy the Windows .env to Pi or enable a service before hardware validation.

## 0:00–0:35 Problem and honest scope

Presenter: “Nova keeps speech processing and personal memory on this machine. It combines
local voice models with deterministic tasks and reminders. The Pi deployment and GPIO
modules exist, but physical Pi validation remains a separate milestone.”
Show online tools disabled, then disconnect external network access. Keep loopback
available. Do not change network settings for another person's machine without consent.

## 0:35–1:15 Wake activation and local answer

Wait for LISTENING. Say “Nova, explain what a solar eclipse is in one sentence.”
Expected: local Ollama answer, Piper speech, then LISTENING resumes. Wording varies.
Say “Nova” alone, then “What is my name?” promptly within eight seconds.
Expected on a new database: unavailable name response. The window permits one command.
Explain: Whisper transcribes before transcript wake filtering. This is not an acoustic
keyword engine. Wait for LISTENING before each new sentence.

## 1:15–2:15 Explicit memory and a one-minute reminder

Say “Nova, my name is Alex.” Expected: direct saved name acknowledgment (no confirmation).
Say “Nova, my favorite food is pizza.” Expected: a save proposal.
Say “Nova, yes.” Expected: confirmed favorite food.
Say “Nova, remind me to drink water in 1 minute.” Expected: explicit scheduling acknowledgment.
Say “Nova, remember that the project demo is in room seven.” Expected: saved note.
Make clear that explicit creation and legacy name updates can save directly.

## 2:15–3:15 Persistence, tasks and audible reminder

Press Ctrl+C and wait for cleanup. Relaunch with the SAME DATABASE_PATH.
Say “Nova, what is my name?” and “Nova, what is my favorite food?”
Expected: Alex and pizza persist. Say “Nova, show my notes.” Expected: synthetic room note.
Say “Nova, add a task to finish my project.” Then “Nova, show my pending tasks.”
Expected: saved task and its typed ID. Complete the displayed ID, then “Nova, yes.”
Do not assume the ID will be1. Await the water reminder and audible playback.
If timing is delayed by inference/restart, wait rather than asserting delivery occurred.

## 3:15–4:00 Mute and privacy boundary

Press M in the focused Windows console. Show MUTED and speak “Nova, add a task to buy a car.”
Press M to unmute, wait for LISTENING, then ask for pending tasks.
Expected: muted speech created no new task. Microphone mute does not promise silent
reminders. Explain accepted earlier work can finish and SQLite is plaintext.
On a validated Pi, instead cycle the switch and show red/green/blue states. A mock
demonstration must be labeled mock. Explain GPIO does not electrically cut mic power.

## 4:00–5:00 Optional passive observation, only if rehearsed

If the venue/participants consent and the feature passed rehearsal: stop gracefully,
set PASSIVE_MEMORY_ENABLED=true in the same isolated session and relaunch. Say without
wake word “I enjoy listening to music.” Wait for the worker's completion count.
Say “Nova, what do I enjoy doing?” Expected: music observation with unidentified-speaker
qualification, not a confirmed owner fact. Background speech must cause no spoken answer.
If this optional feature is omitted, show its code/diagram path and label it opt-in.
Do not stage a response as if it came from the live microphone.

## 5:00–6:00 Limits and conclusion

Ask “Nova, what is the weather in Delhi?” With online tools disabled expect the explicit
offline-unavailable response. Do not claim a fresh offline forecast. Point to the
architecture boundary and explain external tools are optional. Close with physical Pi
validation, sustained testing and acoustic wake detection as next milestones.
Press Ctrl+C, wait for accepted work to drain, and verify ordinary temporary WAV cleanup.

## Fallbacks and recovery

- If voice fails, switch to scripts/assistant_test.py using the isolated database and
  label the demonstration text mode. It demonstrates routing/persistence, not STT/TTS.
- If the LLM is unavailable, show deterministic recall/tasks. Admit general chat is
  unavailable until local Ollama is healthy. Do not expose raw private logs.
- If a reminder has not spoken, show the pending record and state delivery remains
  unverified. Saved scheduling is different from audible success.
- If GPIO fails, retain mute and repair/restart. Do not bypass fail-closed behavior.
- Keep a rehearsed recording with synthetic data only as a labeled prior recording.
  This audit does not create or claim a physical demo video.
- Skip spoken interruption, live online providers and unvalidated Pi boot behavior
  in the core jury sequence. Present them with their actual status.
''')

write('evaluation-readiness-checklist.md', '''# Final evaluation readiness checklist

Unchecked items require presenter action. This audit does not certify live hardware.

## Content and claims

- [x] Architecture source and exported PNG created, with online/control/privacy boundaries.
- [x] 12-slide PPTX created with notes and source references.
- [x] At least30 jury questions answered using repository evidence.
- [x] Demo script and feature-status audit created.
- [ ] Read verification.md for the new test result and runtime limitations.
- [ ] Rehearse notes and keep the pitch within the event's allocated time.
- [ ] Remove unsupported claims of model invention, encrypted memory or electrical mute.
- [ ] Call the Pi a deployment target until physical evidence exists.
- [ ] Label Windows manual checks as reported, benchmarks as historical and mocks as mocks.
- [ ] Open the final PPTX on the presentation machine and check fonts/notes/projector output.

## Reproducible environment

- [ ] Restore a working supported Python3.13 environment and run its full test suite.
- [ ] Confirm project audio/HTTP/timezone dependencies and local binaries/models.
- [ ] Use scripts/always_on_assistant.py rather than the placeholder packaged CLI.
- [ ] Confirm local Ollama API serves llama3.2:1b with no external model forwarding.
- [ ] Record model versions/digests and actual device specs in the final event evidence.
- [ ] Verify Whisper tiny.en and matching Piper ONNX/config pair before the event.
- [ ] Keep the demo SQLite database separate from production personal memories.
- [ ] Do not publish .env, API keys, private DBs, audio, prompts or terminal content.

## Physical voice demo

- [ ] Check mic input/speaker output and16kHz mono input on the actual device.
- [ ] Rehearse venue noise, mic distance and VAD threshold with synthetic speech.
- [ ] Test wake+command, wake-only follow-up and expired eight-second follow-up.
- [ ] Verify unprefixed speech causes no active spoken response.
- [ ] Verify basic general chat, Piper speech and return to LISTENING.
- [ ] Save and recall a confirmed fact, note and task after restart.
- [ ] Show a declined/expired confirmation leaves data unchanged.
- [ ] Verify unknown memory gives an honest unavailable response.
- [ ] Verify task completion uses the displayed ID and scoped confirmation.
- [ ] Verify one-minute reminder playback, completion and restart recovery.
- [ ] Test input mute during idle/capture/STT/output and rejection of stale input.
- [ ] Verify owned WAV cleanup on normal shutdown and explain abrupt-exit limits.
- [ ] Keep experimental interruption disabled for the core pitch unless separately rehearsed.
- [ ] Keep online tools disabled and demonstrate external-network disconnection.
- [ ] Prepare clearly labeled text-mode and prior-recording fallbacks.

## Optional passive/online demonstration

- [ ] Obtain consent for passive collection and explicitly enable it for synthetic data.
- [ ] Wait for extraction completion before recall and show unidentified-speaker wording.
- [ ] Test duplicates, negation, conflicting observations and confirmed-fact priority.
- [ ] Disable passive collection afterward and allow accepted work to drain.
- [ ] If showing weather/news, configure actual providers and keys privately.
- [ ] Verify live provider quota/key/network failure behavior before the event.
- [ ] Clearly state provider queries cross the Internet boundary.

## Raspberry Pi5 hardware and deployment gate

- [ ] Verify Linux64-bit ARM, Python3.13 and actual ARM64 package/wheel compatibility.
- [ ] Provision explicitly, verify trusted Ollama digest and local model presence.
- [ ] Verify ALSA/PortAudio devices as the non-root service user.
- [ ] Inspect gpiozero2/lgpio API compatibility, RP1 chip index and audio/gpio group access.
- [ ] Verify BCM pin mapping, LED polarity, per-color resistors and switch wiring.
- [ ] Test red mute, green listening, blue processing/speaking and shutdown off.
- [ ] Verify every restart starts muted even if the switch already indicates listen.
- [ ] Verify stable mute-to-listen cycle, bounce handling and physical mute priority.
- [ ] Simulate supported GPIO read/write errors: latch mute until explicit restart.
- [ ] Explain software mute does not electrically disconnect microphone power/data.
- [ ] Run target systemd-analyze verify before service installation.
- [ ] Enable units only after the manual physical checklist passes.
- [ ] Test reboot/start ordering, exit75 retry, exit78 stop and failure rate limit.
- [ ] Test SIGTERM during capture/STT/LLM/playback and accepted passive drain.
- [ ] Verify LED/pin/child/audio cleanup and SQLite consistency after shutdown.
- [ ] Verify service stop can be investigated if infinite draining stalls.
- [ ] Run the deferred multi-hour physical soak with CPU/RSS/temperature and latency logs.

## Unfinished and unverified scope to disclose

- [ ] Physical Pi GPIO, headless audio, ARM64 inference and boot recovery remain unverified.
- [ ] Multi-hour physical reliability and target thermal/resource capacity remain unverified.
- [ ] Live online-provider/key/quota tests remain unperformed in recorded evidence.
- [ ] No acoustic wake engine, AEC, general full-duplex or authenticated speaker identity.
- [ ] No encrypted SQLite, secure erasure or electrically enforced mic cutoff.
- [ ] Whole-model-context token budgeting and large/multi-user archive performance incomplete.
- [ ] Root README/security/license/setup docs and portable packaging need follow-up.
''')

# Explicitly requested architecture drawing: editable SVG plus matching PNG.
W,H = 2000,1200
im = Image.new('RGB', (W,H), '#FFFFFF')
draw = ImageDraw.Draw(im)
svg = [f'<svg xmlns="http://www.w3.org/2000/svg" width="{W}" height="{H}" viewBox="0 0 {W} {H}">',
       '<rect width="2000" height="1200" fill="white"/>']
FONT_DIR = Path('C:/Windows/Fonts')
def font(size,bold=False):
    return ImageFont.truetype(str(FONT_DIR / ('arialbd.ttf' if bold else 'arial.ttf')),size)
def txt(x,y,s,size=24,color=INK,bold=False):
    draw.text((x,y),s,font=font(size,bold),fill=color)
    svg.append(f'<text x="{x}" y="{y+size*.86}" font-family="Arial, sans-serif" font-size="{size}" fill="{color}" font-weight="{700 if bold else 400}">{escape(s)}</text>')
def rect(x,y,w,h,fill='white',stroke='#CADADD',width=2,dashed=False):
    draw.rectangle((x,y,x+w,y+h),fill=fill,outline=stroke,width=width)
    svg.append(f'<rect x="{x}" y="{y}" width="{w}" height="{h}" fill="{fill}" stroke="{stroke}" stroke-width="{width}"'+(' stroke-dasharray="10 6"' if dashed else '')+'/>')
def box(x,y,w,h,title,body,fill=PALE):
    rect(x,y,w,h,fill)
    txt(x+16,y+14,title,26,INK,True)
    for i,line in enumerate(body): txt(x+16,y+52+30*i,line,22)
def arrow(points,color=TEAL,dashed=False,label=None,label_xy=None):
    if dashed:
        for p,q in zip(points,points[1:]):
            length=((q[0]-p[0])**2+(q[1]-p[1])**2)**.5
            for n in range(0,int(length),16):
                end=min(n+9,length)
                draw.line((p[0]+(q[0]-p[0])*n/length,p[1]+(q[1]-p[1])*n/length,
                           p[0]+(q[0]-p[0])*end/length,p[1]+(q[1]-p[1])*end/length),fill=color,width=3)
    else: draw.line(points,fill=color,width=3)
    p,q=points[-2:]; dx=q[0]-p[0];dy=q[1]-p[1]; length=(dx*dx+dy*dy)**.5
    ux,uy=dx/length,dy/length
    tri=[q,(q[0]-14*ux+7*uy,q[1]-14*uy-7*ux),(q[0]-14*ux-7*uy,q[1]-14*uy+7*ux)]
    draw.polygon(tri,fill=color)
    svg.append(f'<polyline points="'+ ' '.join(f'{x},{y}' for x,y in points) +f'" fill="none" stroke="{color}" stroke-width="3"'+(' stroke-dasharray="9 7"' if dashed else '')+'/>')
    svg.append('<polygon points="'+' '.join(f'{x},{y}' for x,y in tri)+f'" fill="{color}"/>')
    if label: txt(*label_xy,label,19,color)

txt(40,25,'NOVA  /  Repository architecture',40,NAVY,True)
txt(40,82,'Implemented modules, optional branches and unverified physical deployment',25,MUTED)
rect(350,145,1270,855,'#F7FBFC','#8FAFB5',2)
txt(375,165,'LOCAL DEVICE  /  Windows prototype or Linux ARM64 target',25,TEAL,True)
txt(40,165,'HARDWARE',24,MUTED,True)
rect(1660,145,300,855,'#FFFAF3','#D6B98D',2,True)
txt(1680,170,'OPTIONAL INTERNET',23,AMBER,True)

box(40,240,260,140,'Microphone',['PortAudio input','16 kHz mono'])
box(390,240,320,140,'Capture + RMS VAD',['Bounded queue + pre-roll','PCM16 session WAV'])
box(770,240,320,140,'Whisper STT',['Local whisper.cpp','tiny.en default'])
box(1150,240,420,140,'Activation + filtering',['Transcript-prefix Nova / hey Nova','One 8-second follow-up'])
arrow([(300,310),(390,310)])
arrow([(710,310),(770,310)],label='WAV',label_xy=(716,281))
arrow([(1090,310),(1150,310)],label='text',label_xy=(1096,281))

box(40,470,260,140,'Speaker',['Local WAV playback','sounddevice'])
box(390,470,320,140,'Piper TTS',['Local ONNX voice','Output WAV + cleanup'])
box(770,470,320,140,'Assistant + tools',['Recall / facts / tasks','Reminders / chat routing'])
box(1150,470,420,140,'Passive worker (opt-in)',['Grounded evidence + deduplication','Unidentified-speaker observations'])
arrow([(1360,380),(1360,415),(930,415),(930,470)],label='active command',label_xy=(970,385))
arrow([(1450,380),(1450,470)],label='background',label_xy=(1465,385))
arrow([(770,540),(710,540)],label='answer',label_xy=(712,504))
arrow([(390,540),(300,540)])

box(1680,470,260,180,'Weather / news',['Fixed HTTPS providers','Open-Meteo / NewsAPI','No memory or history','OFF by default'],'#FFF0DB')
arrow([(1090,505),(1120,505),(1120,445),(1640,445),(1640,530),(1680,530)],AMBER,True,label='explicit location / configured topic only',label_xy=(1200,416))
arrow([(1680,600),(1640,600),(1640,690),(740,690),(740,580),(710,580)],AMBER,True)
txt(1190,662,'Sanitized provider answer bypasses Ollama',20,AMBER)

box(390,745,320,140,'Ollama Llama 3.2 1B',['Loopback /api/generate','Chat + candidate extraction'])
box(770,745,320,140,'SQLite + RAM history',['Facts, notes, tasks, reminders','Evidence / times / RAM turns'])
box(1150,745,420,140,'Reminder scheduler',['UTC rows + delivery leases','Playback acknowledgment + retry'])
arrow([(825,610),(825,640),(550,640),(550,745)],label='local prompt',label_xy=(428,655))
arrow([(710,805),(745,805),(745,565),(770,565)],label='model reply',label_xy=(596,713))
arrow([(975,610),(975,745)],label='read / write',label_xy=(985,703))
arrow([(1280,610),(1280,710),(1050,710),(1050,745)],label='validated writes',label_xy=(1100,715))
arrow([(1150,815),(1090,815)])
arrow([(1360,885),(1360,950),(740,950),(740,670),(550,670),(550,610)],dashed=True,label='notify main loop; reminder playback through Piper',label_xy=(770,958))
arrow([(1150,580),(1120,580),(1120,725),(720,725),(720,850),(710,850)],label='optional extraction',label_xy=(765,698))

box(40,745,260,140,'Mute controls',['GPIO switch / console','BCM23 default'])
box(40,920,260,80,'RGB LED',['R17 / G27 / B22'])
arrow([(300,800),(335,800),(335,1030),(460,1030)],NAVY,True)
arrow([(770,1030),(335,1030),(335,960),(300,960)],NAVY,True)
rect(390,1015,1550,135,'#EAF1F8','#8CA4BB',2)
txt(410,1033,'PRIVACY + CONTROL',24,NAVY,True)
txt(410,1070,'Mute generations deny new/stale input. Accepted work can drain. GPIO faults latch mute until restart.',23)
txt(410,1107,'Half-duplex capture/playback. Owned WAV cleanup. Plaintext SQLite. No electrical mic cutoff or voice authentication.',22)
arrow([(570,1015),(570,995),(370,995),(370,355),(390,355)],NAVY,True)
txt(40,1162,'Solid teal: local data     Dashed navy: control/state     Dashed amber: optional network     Pi hardware / boot / inference: unverified',22,MUTED)
svg.append('</svg>')
(ARCH/'nova-architecture.svg').write_text('\n'.join(svg),encoding='utf-8')
im.save(ARCH/'nova-architecture.png')
shutil.copy2(ARCH/'nova-architecture.svg',OUT/'nova-architecture.svg')
shutil.copy2(ARCH/'nova-architecture.png',OUT/'nova-architecture.png')
(ARCH/'README.md').write_text('''# Nova architecture

nova-architecture.svg is the editable diagram source. nova-architecture.png is the
exported image, also embedded in the hackathon PPTX. Copies live in
docs/hackathon-presentation to keep the evaluation package self-contained.

Solid teal arrows show local data. Dashed navy arrows show control/state.
Dashed amber arrows show optional network requests/responses. The privacy strip
applies across capture, STT, activation and storage. The main loop handles active
tools and reminder playback. GPIO observes state and supplies software permission.
The Pi5 is a target: physical GPIO, ARM64 inference and boot behavior remain unverified.
The pipeline is sequential and half-duplex, with optional experimental stop monitoring.

Read ../hackathon-presentation/repository-evaluation.md for full flow, evidence,
configuration defaults, qualifications and privacy limitations.
''',encoding='utf-8')

# A readable 16:9 deck with native editable text/tables and embedded evidence image.
prs = Presentation()
prs.slide_width,prs.slide_height = Inches(13.333),Inches(7.5)
prs.core_properties.title = 'Nova: Personal Offline Always-On AI Companion'
prs.core_properties.subject = 'Repository-grounded hackathon evaluation'
prs.core_properties.author = 'Nova project'
prs.core_properties.keywords = 'Nova, offline, architecture, hackathon, evidence'

def rgb(color): return RGBColor.from_string(color.lstrip('#'))
def text(slide,x,y,w,h,s,size=24,color=INK,bold=False):
    shape = slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    tf=shape.text_frame; tf.clear(); tf.word_wrap=True
    tf.margin_left=tf.margin_right=0
    tf.margin_top=tf.margin_bottom=0
    for i,line in enumerate(s.split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.text=line; p.font.name='Arial';p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=rgb(color)
        p.space_after=Pt(10)
    return shape
def slide(title,notes):
    s=prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid();s.background.fill.fore_color.rgb=rgb('#FFFFFF')
    text(s,.6,.32,12,.3,'NOVA',12,TEAL,True)
    text(s,.6,.85,12.1,.8,title,34,NAVY,True)
    text(s,12.15,7.03,.5,.25,str(len(prs.slides)).zfill(2),11,MUTED)
    s.notes_slide.notes_text_frame.text=notes
    return s
def col(s,x,heading,body):
    text(s,x,2.05,5.65,.55,heading,23,TEAL,True)
    text(s,x,2.82,5.65,3.8,body,23,INK)
def table(s,rows,widths,y=2.05):
    t=s.shapes.add_table(len(rows),len(rows[0]),Inches(.6),Inches(y),Inches(12.1),Inches(4.45)).table
    for i,w in enumerate(widths):t.columns[i].width=Inches(w)
    for r,row in enumerate(rows):
        t.rows[r].height=Inches(.61 if r==0 else .76)
        for c,value in enumerate(row):
            cell=t.cell(r,c);cell.text=value;cell.margin_left=Inches(.12);cell.margin_top=Inches(.1)
            cell.fill.solid();cell.fill.fore_color.rgb=rgb(NAVY if r==0 else ('#F0F6F7' if r%2 else '#FFFFFF'))
            for p in cell.text_frame.paragraphs:
                p.font.name='Arial';p.font.size=Pt(17);p.font.bold=r==0;p.font.color.rgb=rgb('#FFFFFF' if r==0 else INK)
    return t

notes=[]
def add_note(title,body,evidence):
    n=f'{body}\n\nRepository evidence: {evidence}\nAudit snapshot: {DATE}. Historical results and unverified hardware are identified explicitly.'
    notes.append((title,n));return n

s=slide('Personal Offline Always-On AI Companion',add_note('Nova',
 'Opening, 25 seconds: Introduce Nova as a local voice companion with persistent personal memory. Explain that the implementation combines established models rather than training a new one. The Windows prototype and automated policies have evidence. Pi5 is a deployment target with implemented preparation and GPIO code, but its physical operation remains unverified. Do not use always-on to imply uninterrupted full-duplex recording.',
 'scripts/always_on_assistant.py; config.py; docs/DEVELOPMENT_STATUS.md'))
text(s,.6,2.15,11.9,1.3,'Speech and personal memory\nprocessed on your own device',36,TEAL,True)
text(s,.6,4.05,10.6,1.25,'Local voice models. Persistent SQLite memory.\nExplicit tools and visible microphone controls.',25,INK)
text(s,.6,6.22,11.8,.4,'Hackathon prototype evaluation  /  10 October 2026',17,MUTED)

s=slide('Problem statement',add_note('Problem statement',
 '35 seconds: Frame these as design motivations, not validated market statistics. Personal assistants should remain useful when the external network is unavailable. Repeatedly restating personal information is inconvenient. Continuous sensing needs understandable controls. Nova attempts to address these needs locally. The repository does not provide user research or quantified market validation.',
 'project description in pyproject.toml; core/privacy.py; memory/database.py'))
col(s,.6,'Everyday need','A voice interface that stays useful\nwithout an Internet connection\n\nSaved preferences, notes and tasks\nthat survive a restart')
col(s,7,'Design requirement','Personal context stays local\nby default\n\nMicrophone permission is visible\nand can be revoked')

s=slide('Proposed solution and uniqueness',add_note('Solution and uniqueness',
 '40 seconds: Explain the complete local integration. Emphasize confirmed facts versus unidentified passive observations, deterministic tools, transactional memory changes and shared privacy gates. Optional passive extraction defaults off. This is integration innovation rather than a new foundation model. Online weather/news are a clearly separate opt-in branch.',
 'core/assistant.py; memory/passive_processor.py; memory/management.py; tools/lifecycle.py'))
col(s,.6,'What Nova does','Wake-prefixed voice commands\nLocal conversation and recall\nPersistent notes, tasks and reminders\nOptional passive observations')
col(s,7,'What makes it distinctive','Confirmed facts outrank observations\nEvidence grounds personal recall\nDeterministic tools control mutations\nMute gates reject stale input')

s=slide('System architecture',add_note('System architecture',
 '60 seconds: Walk the top row from mic through RMS VAD to whisper.cpp and transcript activation. Follow active commands to Assistant, local memory or loopback Ollama, then Piper and speaker. Background speech enters only the optional passive worker. Explain reminders route through the main output loop. The amber boundary is explicit weather/news, not cloud inference. Navy dashed arrows are control. GPIO supplies software permission and reflects state. Plaintext storage and unverified Pi hardware are visible limitations.',
 '; '.join(anchors)))
s.shapes.add_picture(str(ARCH/'nova-architecture.png'),Inches(.6),Inches(1.72),width=Inches(12.1),height=Inches(5.22))

s=slide('Technology stack',add_note('Technology stack',
 '35 seconds: All stack choices come from adapters and configuration. Clarify that the declared application runtime is Python3.13, while bundled Python3.12 is used only for artifact authoring and qualified audit checks. tiny.en and lessac are platform defaults. GPIO uses local gpiozero/LGPIOFactory rather than remote pins. No embeddings are in this implementation.',
 'pyproject.toml; config.py; core/platform_config.py; core/hardware.py; deployment/raspberry-pi/nova.env.example'))
table(s,[['Layer','Repository choice','Purpose'],['Audio','sounddevice / soundfile / NumPy','Capture, RMS VAD and playback'],['Speech','whisper.cpp tiny.en / Piper lessac','Local speech recognition and synthesis'],['Reasoning','Ollama / llama3.2:1b','Loopback, non-streaming generation'],['Memory','SQLite / concept ranking','Persistent records and grounded recall'],['Runtime','Python3.13 / Linux ARM64 / GPIO','Orchestration and Pi deployment target']],[1.8,5.15,5.15])

s=slide('Working flow',add_note('Working flow',
 '40 seconds: A typical utterance ends after one second of silence. STT precedes wake handling. Wake+command activates directly, and wake-only permits a single eight-second follow-up. Most tool/recall paths avoid general model generation. Capture resumes after output and0.3-second settling. Background speech does not receive an answer. The flow is half-duplex, so users should wait for listening before speaking again.',
 'audio/continuous_recorder.py; stt/whisper.py; core/activation_controller.py; scripts/always_on_assistant.py'))
text(s,.6,2.05,11.9,3.9,'01  Capture a bounded utterance and run local STT\n02  Match Nova at the transcript prefix\n03  Resolve grounded recall, a tool or local chat\n04  Generate local speech and play the answer\n05  Clean owned audio and resume listening',25,INK)
text(s,.6,6.2,12,.45,'Wake-only follow-up: 8 seconds. Normal operation is sequential and half-duplex.',18,TEAL,True)

s=slide('Implementation and memory integrity',add_note('Implementation and memory integrity',
 '45 seconds: SQLite holds separate facts, notes, tasks, reminders and evidence-bearing passive records. Conversation history stays in RAM with ten-turn and4,000-character bounds. Hybrid recall uses a bounded concept lexicon, not embeddings. Management confirmations expire and recheck rows within a transaction. Do not say every creation requires confirmation: legacy name, explicit note/task and reminder creation are direct-write exceptions. Historical1,278-pass checkpoint is not the new audit result. Read verification.md before presenting updated checks.',
 'memory/database.py; memory/hybrid_retriever.py; core/conversation.py; memory/management.py; tools/lifecycle.py; docs/DEVELOPMENT_STATUS.md'))
col(s,.6,'Implemented safeguards','Separate confirmed and passive memory\nEvidence, timestamps and deduplication\nScoped corrections and deletions\nReminder leases and acknowledgments')
col(s,7,'Verification boundary','Historical checkpoint: 1,278 tests passed\nVoice/GPIO boundaries often use mocks\nPi hardware and sustained run pending\nFresh audit result in verification.md')

s=slide('Hardware integration and Pi deployment',add_note('Hardware and Pi deployment',
 '45 seconds: Explain intended wiring and fail-closed permission behavior, not a proven hardware demo. Default BCM17/27/22 drive RGB and BCM23 reads the switch. Listening green, processing/speaking blue and mute red. Restart starts muted and needs deliberate stable switch cycling. GPIO faults latch mute. Headless production requires LinuxARM64 with local lgpio. Provisioning/setup and systemd templates exist, but actual GPIO, ALSA, boot and model performance are unverified. The switch does not disconnect mic power.',
 'core/hardware.py; core/service_runtime.py; core/pi_deployment.py; deployment/raspberry-pi/README.md'))
col(s,.6,'GPIO implementation','RGB: BCM17 / 27 / 22\nMute input: BCM23\n20 ms polling, 100 ms release debounce\nFaults latch mute until restart')
col(s,7,'Deployment target','Raspberry Pi5, Linux64-bit ARM\nPython3.13 and local models\nGuarded launch and non-root systemd\nPhysical validation remains pending')
text(s,.6,6.3,12,.45,'The switch controls software permission. It does not electrically disconnect the microphone.',17,AMBER,True)

s=slide('Live demonstration',add_note('Live demonstration',
 'Demonstration, about two minutes within the pitch if time permits: Use a synthetic isolated SQLite path and preloaded models. Say Nova explain a solar eclipse in one sentence. Save Alex via name command and pizza via proposed favorite+yes. Add a one-minute water reminder. Restart with same DB and recall name/favorite. Show a task and its displayedID. Mute withM and verify muted speech creates no task. Explain exact wording can vary. Use the actual always_on_assistant script, not the startup-only packagedCLI. If the interpreter/device fails, label the text-mode fallback honestly.',
 'scripts/always_on_assistant.py; scripts/assistant_test.py; demo-script.md; tools/memory_tools.py'))
col(s,.6,'Core jury sequence','Local question with network disconnected\nSave a synthetic preference\nRestart and recall the same fact\nCreate a task and deliver a reminder')
col(s,7,'Visible privacy control','Mute input and show the state\nVerify muted speech creates no task\nOptional passive path only with consent\nLabel text-mode and mock fallbacks')

s=slide('Challenges and current limitations',add_note('Challenges and limitations',
 '45 seconds: Disclose limits concretely. RMSVAD is susceptible to venue noise and transcript-based activation incursSTT cost before wake. Sequential processing has capture gaps. Small model context can constrain memory prompts. SQLite is plaintext, speaker identity is absent and electrical cutoff is absent. NoAEC/full-duplex exists. Pi, provider and multi-hour physical evidence need validation. The audit found StorePython/venv inaccessible in this environment, empty root docs and a nonportable Windows freeze. These must be addressed before claiming easy deployment.',
 'audio/vad.py; core/wake_detector.py; core/hardware.py; core/assistant.py; docs/DEVELOPMENT_STATUS.md; requirements.txt'))
col(s,.6,'Technical constraints','Noise-sensitive energy VAD\nSTT runs before wake activation\nSequential capture and model latency\nSmall context and shared CPU budget')
col(s,7,'Evaluation limits','Physical Pi and multi-hour run pending\nPlaintext memory and no speaker identity\nNo AEC or electrical microphone cutoff\nPackaging and public docs incomplete')

s=slide('Future scope',add_note('Future scope',
 '35 seconds: These are recommendations, not completed features. Prioritize physical Pi validation and repeatable packaging before expanding scope. Measure thermal/resource/latency and long-running reliability. Then consider acoustic wake/VAD, authenticated voice commands, consent/retention controls, encrypted storage and true electrical mic disconnect. Full-duplex and broader integrations require their own implementation. No invented milestones, user metrics or promised model performance.',
 'repository-evaluation.md recommendations; docs/LONG_RUNNING_VALIDATION.md; deployment/raspberry-pi/README.md'))
col(s,.6,'Next validation milestone','Real Pi installation and boot recovery\nHeadless audio and GPIO fault checks\nMulti-hour thermal/resource measurements\nPortable setup and voice entry point')
col(s,7,'Next product capabilities','Dedicated acoustic wake detection\nBetter noise handling and echo control\nAuthenticated actions and encrypted memory\nToken-bounded context and archive tests')

s=slide('Conclusion',add_note('Conclusion',
 'Closing,25 seconds: Nova demonstrates a coherent local voice architecture with persistent memory, deterministic tools and software privacy gates. The repository contains substantial automated test coverage and bounded recovery policies. State clearly that target hardware and sustained physical reliability are next validation gates. Invite questions and keep jury-cross-questions.md nearby. Avoid promising encrypted, authenticated or uninterrupted listening behavior.',
 'scripts/always_on_assistant.py; docs/DEVELOPMENT_STATUS.md; repository-evaluation.md'))
text(s,.6,2.1,11.9,1.5,'Local voice assistance with\npersistent, evidence-aware memory',35,TEAL,True)
text(s,.6,4.08,11.9,1.2,'Implemented orchestration, tools and privacy controls\nA clear path to physical Pi validation',25,INK)
text(s,.6,6.22,11.8,.4,'Repository evidence defines the claims. Live hardware defines the next milestone.',18,MUTED)

PPTX = OUT / 'nova-hackathon-presentation.pptx'
prs.save(PPTX)
write('speaker-notes.md','# Speaker notes\n\n'+'\n\n'.join(f'## {i}. {title}\n\n{note}' for i,(title,note) in enumerate(notes,1)))

# Verify package, XML, slide count, notes, native tables, embedded PNG and geometry.
checks=[]
with ZipFile(PPTX) as z:
    checks.append({'check':'zip_crc','ok':z.testzip() is None})
    xmls=[n for n in z.namelist() if n.endswith(('.xml','.rels'))]
    for n in xmls:ET.fromstring(z.read(n))
    checks.append({'check':'all_xml_parse','ok':True,'count':len(xmls)})
    checks.append({'check':'12_slides','ok':len(prs.slides)==12})
    ns=[n for n in z.namelist() if n.startswith('ppt/notesSlides/notesSlide') and n.endswith('.xml')]
    checks.append({'check':'12_notes','ok':len(ns)==12})
    checks.append({'check':'architecture_png_embedded','ok':any(n.startswith('ppt/media/') for n in z.namelist())})
reloaded=Presentation(PPTX)
for i,s in enumerate(reloaded.slides,1):
    assert s.notes_slide.notes_text_frame.text.strip()
    for sh in s.shapes:
        assert sh.left>=0 and sh.top>=0 and sh.left+sh.width<=prs.slide_width and sh.top+sh.height<=prs.slide_height, (i,sh.name)
checks.append({'check':'all_shapes_within_canvas','ok':True})
checks.append({'check':'jury_questions_at_least_30','ok':len(QA)>=30,'count':len(QA)})
ET.parse(ARCH/'nova-architecture.svg')
Image.open(ARCH/'nova-architecture.png').verify()
checks.append({'check':'diagram_svg_and_png_parse','ok':True})
checks.append({'check':'source_and_tests_syntax_parsed','ok':True,'count':sum('syntax' in x for x in inventory)})
write('artifact-validation.json',json.dumps(checks,indent=2))
write('README.md','''# Nova hackathon evaluation package

- nova-hackathon-presentation.pptx: 12 slides with embedded architecture and speaker notes.
- repository-evaluation.md: architecture, feature matrix, evidence, limitations and risks.
- jury-cross-questions.md: 48 concise repository-grounded jury questions and answers.
- demo-script.md: six-minute demonstration with synthetic data and labeled fallbacks.
- evaluation-readiness-checklist.md: presenter, hardware and deployment gates.
- speaker-notes.md: standalone copy of slide notes.
- nova-architecture.svg / .png: self-contained diagram copies.
- source-inventory.json: first-party text-file hashes and Python symbol inventory.
- artifact-validation.json / verification.md: generated-file and test verification.
- build_deliverables.py: reproducible artifact builder (existing python-pptx and Pillow).

Canonical architecture files also live in ../architecture/ as requested.
Slide previews and test-results.xml, when present, document the separate verification
step. No application source changes or dependency installations are part of this package.
Private .env, personal memory contents and recordings are excluded.
''')
print(json.dumps({'slides':len(prs.slides),'questions':len(QA),'inventory_files':len(inventory),'checks':checks},indent=2))
