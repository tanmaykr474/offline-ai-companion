from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = Path(__file__).parent / 'Nova Evaluation Commands.docx'
d = Document()
s = d.sections[0]
s.top_margin = s.bottom_margin = Inches(.65)
s.left_margin = s.right_margin = Inches(.7)
d.styles['Normal'].font.name = 'Calibri'
d.styles['Normal'].font.size = Pt(10)
d.styles['Normal'].paragraph_format.space_after = Pt(5)
for name in ['Title', 'Heading 1', 'Heading 2']:
    d.styles[name].font.color.rgb = RGBColor(0, 0, 0)

def p(text): d.add_paragraph(text)
def table(rows):
    t = d.add_table(rows=1, cols=3)
    t.style = 'Light Shading Accent 1'
    for c, text in zip(t.rows[0].cells, ['Feature', 'Say or do', 'Expected result']): c.text = text
    repeat = OxmlElement('w:tblHeader'); t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in rows:
        cells = t.add_row().cells
        for c, text in zip(cells, row): c.text = text
        keep = OxmlElement('w:cantSplit'); t.rows[-1]._tr.get_or_add_trPr().append(keep)
    for row in t.rows:
        for c, width in zip(row.cells, [1.1, 3.05, 2.65]): c.width = Inches(width)
    return t
def page(title):
    d.add_page_break(); d.add_heading(title, 1)

d.add_paragraph('Nova Evaluation Commands', 'Title')
p('Windows live demonstration and optional feature checks')
p('Use the validated Windows setup for the evaluation. Say each command separately and wait for LISTENING before the next command. Prefix confirmations with Nova. This checklist gives expected behavior; it is not a record that every scenario has passed.')
d.add_heading('Recommended five minute demonstration', 1)
table([
('Greeting', 'Nova, hello.', 'A spoken response, followed by listening.'),
('Save a fact', 'Nova, my favorite food is biryani.\nNova, yes. (If asked.)', 'Acknowledges the saved fact.'),
('Recall', 'Nova, what is my favorite food?', 'Answers biryani.'),
('Paraphrase', 'Nova, what food do I like most?', 'Retrieves the same saved fact; rehearse first.'),
('Note', 'Nova, remember that the project demo is in room seven.\nNova, show my notes.', 'Saves and lists the note.'),
('Task', 'Nova, add a task to finish my project.\nNova, show my pending tasks.', 'Lists the task with its stable ID.'),
('Reminder', 'Nova, remind me to drink water in 1 minute.', 'Acknowledges scheduling; announces once when due.'),
('Mute', 'Press M in the focused Windows console.\nSay Nova, hello.\nPress M again; repeat the greeting.', 'No activation while muted; responds after unmute.'),
('Persistence', 'Ctrl+C, restart Nova, then:\nNova, what is my favorite food?', 'Saved fact survives restart.')])
p('Do not blindly repeat a save or scheduling command after a missing spoken response. Check the relevant list first. Record one successful demonstration as a backup.')

page('Personal memory and notes')
table([
('Name', 'Nova, my name is Tanmay.\nNova, what is my name?', 'Saves and recalls the name.'),
('Correction', 'Nova, my favorite food is now pizza.\nNova, yes.\nNova, what is my favorite food?', 'Confirmed answer becomes pizza.'),
('Cancel change', 'Nova, my favorite food is now pasta.\nNova, cancel.\nNova, what is my favorite food?', 'Previous confirmed value remains.'),
('Review', 'Nova, show my saved memories.\nNova, list my preferences.\nNova, show my saved facts.', 'Bounded review distinguishing record categories.'),
('Pagination', 'Nova, show my saved memories page 2.', 'Next page, or an honest empty-page response.'),
('Passive review', 'Nova, show my passive memories.', 'Nonexpired observations with speaker attribution.'),
('Inspect record', 'Nova, show memory number 5.', 'Passive record 5 if it exists. Use an actual listed ID.'),
('Delete preference', 'Nova, remove my favorite food preference.\nNova, cancel.', 'Proposes deletion; cancellation leaves data intact.'),
('Delete observation', 'Nova, delete memory number 5.\nNova, yes.', 'Deletes only the selected passive record. Use disposable data.'),
('Bulk safety', 'Nova, clear my passive memories.\nNova, cancel.', 'States scope; cancellation deletes nothing.'),
('Other bulk scopes', 'Nova, delete all saved facts.\nNova, cancel.\nNova, forget everything you remember about me.\nNova, cancel.', 'Confirmed facts and passive observations have explicit scopes; notes/tasks/reminders are separate.'),
('Missing fact', 'Ask about a personal preference you have never saved.', 'Does not invent personal information.')])
p('Memory numbers identify passive observations, not notes, tasks or confirmed facts. Deleting a confirmed preference can leave a separate passive observation available. Never bulk-delete your evaluation data just to demonstrate the feature.')

page('Tasks and reminders')
p('Replace example IDs with IDs returned by Nova. Complete, rename, delete, cancel, snooze and reschedule require confirmation. Use Nova, yes or Nova, cancel.')
table([
('Rename task', 'Nova, rename my project task to finish the API integration.\nNova, yes.', 'Renames a uniquely matched task.'),
('Complete task', 'Nova, complete task number 3.\nNova, yes.\nNova, show my pending tasks.', 'Selected task disappears from pending tasks.'),
('Repeat completion', 'Nova, complete task number 3.', 'Already-completed task is not changed again.'),
('Delete task', 'Nova, delete task number 3.\nNova, yes.', 'Deletes only the selected task.'),
('Upcoming reminders', 'Nova, show my upcoming reminders.', 'Lists active reminders with IDs.'),
('Tomorrow', 'Nova, remind me to call Mom tomorrow at 6 PM.', 'States resolved date, time and configured timezone.'),
('Absolute date', 'Nova, remind me about the meeting on October 15 2026 at 10 AM.', 'Schedules only if still future. Use a future date when testing later.'),
('Missing message', 'Nova, remind me in 30 minutes.', 'Requests the message; does not invent one.'),
('Snooze', 'Nova, snooze reminder number 2 for 10 minutes.\nNova, yes.', 'Moves that reminder to a later time.'),
('Reschedule', 'Nova, reschedule my meeting reminder to tomorrow at 11 AM.\nNova, yes.', 'Updates the uniquely matched reminder.'),
('Cancel reminder', 'Nova, cancel my reminder to call Mom.\nNova, yes.', 'Makes the selected reminder inactive.'),
('Delete reminder', 'Nova, delete reminder number 2.\nNova, yes.', 'Removes only that reminder.')])
p('Calendar reminders use REMINDER_TIMEZONE, normally Asia/Kolkata. Use AM or PM. Past times, ambiguous numeric dates and unsupported date forms should be rejected. Multiple matching records require a specific ID.')

page('Optional memory and online tests')
d.add_heading('Passive memory', 2)
p('Passive memory must be enabled. Speak the first sentence without Nova, wait for a successful save, then ask the question. The recent live attempt saved zero observations, so rehearse these before evaluation. Unknown speakers must remain attributed as observations.')
table([
('Activity', 'I love playing cricket.\nThen: Nova, what sport do I love playing?', 'Grounded cricket observation, if saved.'),
('Birthplace', 'I was born in Bihar.\nThen: Nova, where was I born?', 'Grounded Bihar observation, if saved.'),
('Enjoyment', 'I enjoy listening to music.\nThen: Nova, what do I enjoy doing?', 'Grounded music observation, if saved.'),
('Occupation', 'I work as a software developer.\nThen: Nova, what do I do for work?', 'Grounded occupation observation, if saved.'),
('Multiple facts', 'Nova, what do you remember about my work and hobbies?', 'Only relevant saved evidence; no invented details.'),
('Negation safety', 'I was not born in Kerala.\nThen: Nova, show my passive memories.', 'No affirmative Kerala birthplace created.'),
('Duplicate safety', 'Repeat an identical supported statement, then review passive memories.', 'No identical duplicate record.')])
d.add_heading('Optional weather and news', 2)
table([
('Weather', 'Nova, what is the temperature in Pune, India?', 'Configured provider result, or clear offline response.'),
('Forecast', 'Nova, will it rain tomorrow?', 'Uses configured location or requests one.'),
('News', 'Nova, tell me today\'s top news.', 'Bounded provider headlines with sources, or offline response.'),
('Topic news', 'Nova, give me technology headlines.', 'Configured topic headlines, or offline response.')])
p('Online tools are disabled by default. Open-Meteo and NewsAPI require explicit configuration; NewsAPI needs credentials. Do not claim live API validation from mocked tests. Do not change provider configuration immediately before the demo.')

page('Privacy interruption and presentation boundaries')
table([
('Credential protection', 'Nova, remember that my password is DemoOnly123.', 'Recognized credential storage is refused. Use only this synthetic example.'),
('Local conversation', 'Nova, explain a solar eclipse in one sentence.', 'Ollama response spoken by Piper.'),
('Offline check', 'After provisioning, disconnect internet and repeat greeting and saved-fact recall.', 'Local voice and memory work; online tools cannot supply live data.'),
('Console interruption', 'If experimental interruption is enabled, request a long explanation and press S.', 'Cancels current response and resumes listening.'),
('Spoken interruption', 'If both interruption options are enabled, say Nova, stop during playback.', 'Stops playback; acoustic feedback limitations remain.'),
('Graceful shutdown', 'Press Ctrl+C.', 'Shutdown completes; accepted passive work drains.')])
p('Interruption does not undo already committed notes, tasks or reminders. Software mute is not electrical microphone disconnection. Routine diagnostics should report states and error categories without exposing private speech.')
d.add_heading('Raspberry Pi demonstration wording', 2)
p('We tested the physical Raspberry Pi setup. Speaker playback works, but microphone recordings currently contain static. Today\'s live software demonstration uses Windows. GPIO controls and automatic startup are implemented but still require physical Pi validation.')
p('Show the Pi and connected peripherals honestly as a prototype. Leave unverified GPIO wiring unpowered. Do not claim successful Pi inference, GPIO or systemd operation until physically tested. Multi-hour physical reliability testing remains deferred.')
d.add_heading('If a response is missed', 2)
p('If still LISTENING, repeat the full wake command once. If PROCESSING, wait. If an invalid transcription is reported, speak closer to the microphone in a quiet room. Check lists before repeating commands with side effects. Use the recorded successful demo if needed.')
d.core_properties.title = 'Nova Evaluation Commands'
d.core_properties.subject = 'Voice command checklist for hackathon evaluation'
d.core_properties.author = 'Nova Project'
d.save(OUT)
print(OUT)
